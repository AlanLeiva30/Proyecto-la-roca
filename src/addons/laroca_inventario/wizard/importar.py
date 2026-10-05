import base64
import io

import openpyxl
import xlsxwriter

from odoo import _, fields, models
from odoo.exceptions import UserError
from odoo.tools import float_compare

HOJAS = {
    'Productos': ['Código', 'Nombre', 'Categoría', 'Precio de venta', 'Costo', 'Mínimo por defecto', 'Objetivo por defecto'],
    'Niveles': ['Gasolinera (código)', 'Código producto', 'Stock mínimo', 'Stock objetivo'],
    'Existencias': ['Almacén (código)', 'Código producto', 'Cantidad'],
}
INSTRUCCIONES = [
    'Plantilla de importación - Grupo La Roca',
    '',
    'Hoja "Productos": crea o actualiza productos por su Código. Si la categoría no existe, se crea.',
    'Hoja "Niveles": stock mínimo y objetivo de cada producto en cada gasolinera.',
    'Hoja "Existencias": cantidad REAL que hay hoy en cada almacén (gasolinera o bodega central).',
    '   Se registra como un ajuste de inventario y queda en el historial de movimientos.',
    '',
    'Reglas:',
    ' - No cambies los títulos de las columnas ni los nombres de las hojas.',
    ' - Podés borrar las hojas o filas que no necesites; las celdas vacías no se modifican.',
    ' - Las cantidades no pueden ser negativas y el objetivo no puede ser menor que el mínimo.',
    ' - Si hay errores, NO se importa nada y se indica la hoja y la fila de cada error.',
]


class LarocaImportar(models.TransientModel):
    """Importación masiva desde Excel (.xlsx) del catálogo, los niveles y las existencias."""
    _name = 'laroca.importar'
    _description = 'Importar datos desde Excel'

    archivo = fields.Binary('Archivo Excel (.xlsx)')
    nombre_archivo = fields.Char('Nombre del archivo')
    resultado = fields.Text('Resultado', readonly=True)

    # ------------------------------------------------------------------
    # Plantilla
    # ------------------------------------------------------------------
    def action_descargar_plantilla(self):
        """Genera la plantilla con los datos actuales, lista para editar."""
        salida = io.BytesIO()
        libro = xlsxwriter.Workbook(salida, {'in_memory': True})
        titulo = libro.add_format({'bold': True, 'bg_color': '#0B3D91', 'font_color': '#FFFFFF', 'border': 1})
        hoja = libro.add_worksheet('Instrucciones')
        hoja.set_column(0, 0, 110)
        for fila, linea in enumerate(INSTRUCCIONES):
            hoja.write(fila, 0, linea, titulo if fila == 0 else None)

        productos = self.env['product.template'].search([('is_storable', '=', True)], order='default_code')
        lineas = self.env['laroca.inventario'].search([], order='gasolinera_id, default_code')
        datos = {
            'Productos': [[p.default_code or '', p.name, p.categ_id.name or '', p.list_price, p.standard_price,
                           p.laroca_stock_minimo, p.laroca_stock_objetivo] for p in productos],
            'Niveles': [[l.gasolinera_id.code, l.default_code or '', l.stock_minimo, l.stock_objetivo] for l in lineas],
            'Existencias': [[l.gasolinera_id.code, l.default_code or '', l.stock_actual] for l in lineas],
        }
        for nombre, columnas in HOJAS.items():
            hoja = libro.add_worksheet(nombre)
            for col, encabezado in enumerate(columnas):
                hoja.write(0, col, encabezado, titulo)
                hoja.set_column(col, col, 24 if col else 18)
            for fila, valores in enumerate(datos[nombre], start=1):
                hoja.write_row(fila, 0, valores)
            hoja.freeze_panes(1, 0)
        libro.close()
        adjunto = self.env['ir.attachment'].create({
            'name': 'plantilla_importacion_la_roca.xlsx',
            'datas': base64.b64encode(salida.getvalue()),
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {'type': 'ir.actions.act_url', 'url': f'/web/content/{adjunto.id}?download=true', 'target': 'self'}

    # ------------------------------------------------------------------
    # Lectura y validación
    # ------------------------------------------------------------------
    def _leer_hojas(self):
        if not self.archivo:
            raise UserError(_('Subí el archivo Excel primero.'))
        if self.nombre_archivo and not self.nombre_archivo.lower().endswith('.xlsx'):
            raise UserError(_('El archivo debe ser .xlsx (Excel).'))
        try:
            libro = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.archivo)), read_only=True, data_only=True)
        except Exception as error:
            raise UserError(_('No se pudo leer el archivo Excel: %s', error)) from error
        hojas = {}
        for nombre, columnas in HOJAS.items():
            if nombre not in libro.sheetnames:
                continue
            filas = list(libro[nombre].iter_rows(values_only=True))
            encabezados = [str(c).strip() if c is not None else '' for c in (filas[0] if filas else [])]
            if encabezados[:len(columnas)] != columnas:
                raise UserError(_('La hoja "%(hoja)s" no tiene las columnas esperadas: %(cols)s',
                                  hoja=nombre, cols=', '.join(columnas)))
            ancho = len(columnas)
            hojas[nombre] = [(numero, (tuple(fila) + (None,) * ancho)[:ancho])  # filas cortas se completan
                             for numero, fila in enumerate(filas[1:], start=2)
                             if any(c not in (None, '') for c in fila)]
        if not hojas:
            raise UserError(_('El archivo no tiene ninguna de las hojas: %s', ', '.join(HOJAS)))
        return hojas

    def _numero(self, valor, errores, ubicacion, campo):
        if valor in (None, ''):
            return None
        try:
            numero = float(valor)
        except (TypeError, ValueError):
            errores.append(_('%(ubic)s: "%(campo)s" debe ser un número (se leyó "%(valor)s").',
                             ubic=ubicacion, campo=campo, valor=valor))
            return None
        if numero < 0:
            errores.append(_('%(ubic)s: "%(campo)s" no puede ser negativo.', ubic=ubicacion, campo=campo))
            return None
        return numero

    def action_importar(self):
        self.ensure_one()
        if not self.env.user.has_group('laroca_inventario.group_laroca_admin'):
            raise UserError(_('Solo el administrador puede importar datos.'))
        hojas = self._leer_hojas()
        errores = []
        Producto = self.env['product.template'].with_context(active_test=False)
        Almacen = self.env['stock.warehouse']
        almacenes = {w.code.upper(): w for w in Almacen.search([])}

        # 1) Productos -----------------------------------------------------
        productos_vals = []
        codigos_nuevos = set()
        for fila, valores in hojas.get('Productos', []):
            ubic = _('Productos, fila %s', fila)
            codigo, nombre, categoria = (str(v).strip() if v not in (None, '') else '' for v in valores[:3])
            if not codigo:
                errores.append(_('%s: falta el código.', ubic))
                continue
            existente = Producto.search([('default_code', '=', codigo)], limit=1)
            if not existente and not nombre:
                errores.append(_('%(ubic)s: el producto %(cod)s es nuevo y necesita nombre.', ubic=ubic, cod=codigo))
                continue
            vals = {}
            if nombre:
                vals['name'] = nombre
            for indice, campo, etiqueta in [(3, 'list_price', 'Precio de venta'), (4, 'standard_price', 'Costo'),
                                            (5, 'laroca_stock_minimo', 'Mínimo por defecto'),
                                            (6, 'laroca_stock_objetivo', 'Objetivo por defecto')]:
                numero = self._numero(valores[indice] if len(valores) > indice else None, errores, ubic, etiqueta)
                if numero is not None:
                    vals[campo] = numero
            productos_vals.append((existente, codigo, categoria, vals))
            codigos_nuevos.add(codigo)

        def producto_por_codigo(codigo):
            return self.env['product.product'].search([('default_code', '=ilike', codigo)], limit=1)

        # 2) Niveles y 3) Existencias: solo validación previa ------------------
        niveles, existencias = [], []
        for fila, valores in hojas.get('Niveles', []):
            ubic = _('Niveles, fila %s', fila)
            cod_alm, cod_prod = (str(v).strip().upper() if v else '' for v in valores[:2])
            gasolinera = almacenes.get(cod_alm)
            if not gasolinera or not gasolinera.laroca_es_gasolinera:
                errores.append(_('%(ubic)s: "%(cod)s" no es el código de una gasolinera.', ubic=ubic, cod=cod_alm))
                continue
            if not (cod_prod in {c.upper() for c in codigos_nuevos} or producto_por_codigo(cod_prod)):
                errores.append(_('%(ubic)s: no existe el producto "%(cod)s".', ubic=ubic, cod=cod_prod))
                continue
            minimo = self._numero(valores[2] if len(valores) > 2 else None, errores, ubic, 'Stock mínimo')
            objetivo = self._numero(valores[3] if len(valores) > 3 else None, errores, ubic, 'Stock objetivo')
            if minimo is not None and objetivo and objetivo < minimo:
                errores.append(_('%s: el objetivo no puede ser menor que el mínimo.', ubic))
                continue
            niveles.append((gasolinera, cod_prod, minimo, objetivo))
        for fila, valores in hojas.get('Existencias', []):
            ubic = _('Existencias, fila %s', fila)
            cod_alm, cod_prod = (str(v).strip().upper() if v else '' for v in valores[:2])
            almacen = almacenes.get(cod_alm)
            if not almacen:
                errores.append(_('%(ubic)s: no existe el almacén "%(cod)s".', ubic=ubic, cod=cod_alm))
                continue
            if not (cod_prod in {c.upper() for c in codigos_nuevos} or producto_por_codigo(cod_prod)):
                errores.append(_('%(ubic)s: no existe el producto "%(cod)s".', ubic=ubic, cod=cod_prod))
                continue
            cantidad = self._numero(valores[2] if len(valores) > 2 else None, errores, ubic, 'Cantidad')
            if cantidad is not None:
                existencias.append((almacen, cod_prod, cantidad))

        if errores:
            detalle = '\n'.join(errores[:25]) + (_('\n… y %s errores más.', len(errores) - 25) if len(errores) > 25 else '')
            raise UserError(_('No se importó nada. Corregí estos errores y volvé a intentar:\n\n%s', detalle))

        # Aplicar ------------------------------------------------------------
        creados = actualizados = 0
        for existente, codigo, categoria, vals in productos_vals:
            if categoria:
                categ = self.env['product.category'].search([('name', '=', categoria)], limit=1) \
                    or self.env['product.category'].create({'name': categoria})
                vals['categ_id'] = categ.id
            if existente:
                existente.write(vals)
                actualizados += 1
            else:
                Producto.create(dict(vals, default_code=codigo, type='consu', is_storable=True))
                creados += 1

        Inventario = self.env['laroca.inventario']
        for gasolinera, cod_prod, minimo, objetivo in niveles:
            linea = Inventario.search([('gasolinera_id', '=', gasolinera.id),
                                       ('product_id', '=', producto_por_codigo(cod_prod).id)], limit=1)
            vals = {k: v for k, v in (('stock_minimo', minimo), ('stock_objetivo', objetivo)) if v is not None}
            if 'stock_minimo' in vals and 'stock_objetivo' not in vals and linea.stock_objetivo \
                    and linea.stock_objetivo < vals['stock_minimo']:
                vals['stock_objetivo'] = 0
            linea.write(vals)

        cantidades_gasolineras, ajustes_bodega = {}, 0
        for almacen, cod_prod, cantidad in existencias:
            producto = producto_por_codigo(cod_prod)
            if almacen.laroca_es_gasolinera:
                linea = Inventario.search([('gasolinera_id', '=', almacen.id), ('product_id', '=', producto.id)], limit=1)
                cantidades_gasolineras[linea] = cantidad
            else:
                ajustes_bodega += self._ajustar_almacen(almacen, producto, cantidad)
        if cantidades_gasolineras:
            Inventario._ajustar_existencias(cantidades_gasolineras, _('Importación desde Excel'))

        self.resultado = _(
            'Importación terminada:\n'
            '• Productos creados: %(creados)s · actualizados: %(actualizados)s\n'
            '• Niveles actualizados: %(niveles)s\n'
            '• Existencias registradas: %(existencias)s (gasolineras: %(gas)s, otros almacenes: %(bod)s con cambios)',
            creados=creados, actualizados=actualizados, niveles=len(niveles), existencias=len(existencias),
            gas=len(cantidades_gasolineras), bod=ajustes_bodega)
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'name': _('Importar desde Excel'),
        }

    def _ajustar_almacen(self, almacen, producto, cantidad):
        """Ajuste de inventario en un almacén que no es gasolinera (ej. Bodega Central)."""
        Quant = self.env['stock.quant']
        ubicacion = almacen.lot_stock_id
        actual = sum(Quant.search([('product_id', '=', producto.id), ('location_id', 'child_of', ubicacion.id)]).mapped('quantity'))
        if not float_compare(actual, cantidad, precision_rounding=producto.uom_id.rounding):
            return 0
        quant = Quant._gather(producto, ubicacion, strict=True)[:1] or Quant.create(
            {'product_id': producto.id, 'location_id': ubicacion.id})
        quant.inventory_quantity = quant.quantity + (cantidad - actual)
        quant.with_context(inventory_name=_('Importación desde Excel'))._apply_inventory()
        return 1
