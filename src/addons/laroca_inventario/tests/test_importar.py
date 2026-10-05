import base64
import io

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, new_test_user, tagged


@tagged('post_install', '-at_install', 'laroca')
class TestLarocaImportar(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.g1 = cls.env['stock.warehouse'].create({'name': 'Importa Uno', 'code': 'IMP1', 'laroca_es_gasolinera': True})
        cls.bodega = cls.env.company._laroca_bodega_central()
        cls.admin = new_test_user(cls.env, login='admin_importa', groups='laroca_inventario.group_laroca_admin')
        cls.encargado = new_test_user(cls.env, login='enc_importa', groups='laroca_inventario.group_laroca_encargado',
                                      laroca_gasolinera_ids=[(6, 0, cls.g1.ids)])

    def _libro(self, productos=(), niveles=(), existencias=()):
        from odoo.addons.laroca_inventario.wizard.importar import HOJAS
        libro = openpyxl.Workbook()
        libro.remove(libro.active)
        for nombre, filas in (('Productos', productos), ('Niveles', niveles), ('Existencias', existencias)):
            hoja = libro.create_sheet(nombre)
            hoja.append(HOJAS[nombre])
            for fila in filas:
                hoja.append(list(fila))
        salida = io.BytesIO()
        libro.save(salida)
        return base64.b64encode(salida.getvalue())

    def _importar(self, archivo, user=None):
        wizard = self.env['laroca.importar'].with_user(user or self.admin).create(
            {'archivo': archivo, 'nombre_archivo': 'datos.xlsx'})
        wizard.action_importar()
        return wizard

    def test_plantilla_se_descarga_y_se_reimporta(self):
        wizard = self.env['laroca.importar'].with_user(self.admin).create({})
        accion = wizard.action_descargar_plantilla()
        self.assertIn('/web/content/', accion['url'])
        adjunto = self.env['ir.attachment'].search([('res_model', '=', 'laroca.importar'), ('res_id', '=', wizard.id)])
        libro = openpyxl.load_workbook(io.BytesIO(base64.b64decode(adjunto.datas)))
        self.assertEqual(libro.sheetnames, ['Instrucciones', 'Productos', 'Niveles', 'Existencias'])
        # Reimportar la plantilla sin cambios no rompe nada
        self._importar(adjunto.datas)

    def test_importacion_completa(self):
        archivo = self._libro(
            productos=[('IMP-001', 'Lentes importados', 'Accesorios importados', 9.5, 4, 6, 12)],
            niveles=[('IMP1', 'IMP-001', 8, 16)],
            existencias=[('IMP1', 'IMP-001', 5), (self.bodega.code, 'IMP-001', 40)],
        )
        wizard = self._importar(archivo)
        self.assertIn('Productos creados: 1', wizard.resultado)
        producto = self.env['product.product'].search([('default_code', '=', 'IMP-001')])
        self.assertEqual(producto.categ_id.name, 'Accesorios importados')
        self.assertEqual(producto.standard_price, 4)
        linea = self.env['laroca.inventario'].search([('gasolinera_id', '=', self.g1.id), ('product_id', '=', producto.id)])
        self.assertEqual((linea.stock_minimo, linea.stock_objetivo, linea.stock_actual), (8, 16, 5))
        self.assertEqual(linea.estado, 'bajo')
        self.assertEqual(producto.with_context(warehouse_id=self.bodega.id).qty_available, 40)
        self.assertTrue(self.env['stock.move'].search([('reference', '=', 'Importación desde Excel'), ('product_id', '=', producto.id)]))

    def test_errores_no_importan_nada(self):
        archivo = self._libro(
            productos=[('IMP-002', 'Válido', '', 1, 1, 1, 2), ('IMP-003', '', '', 1, 1, 1, 1)],  # sin nombre
            niveles=[('NOEXISTE', 'IMP-002', 1, 2), ('IMP1', 'IMP-002', 10, 5)],               # gasolinera mala, objetivo < mínimo
            existencias=[('IMP1', 'IMP-002', -4), ('IMP1', 'XYZ', 1)],                           # negativo, producto inexistente
        )
        with self.assertRaises(UserError) as error:
            self._importar(archivo)
        mensaje = str(error.exception)
        for texto in ['Productos, fila 3', 'Niveles, fila 2', 'Niveles, fila 3', 'Existencias, fila 2', 'Existencias, fila 3']:
            self.assertIn(texto, mensaje)
        self.assertFalse(self.env['product.product'].search([('default_code', '=', 'IMP-002')]))

    def test_columnas_incorrectas(self):
        libro = openpyxl.Workbook()
        libro.active.title = 'Productos'
        libro.active.append(['Code', 'Name'])
        salida = io.BytesIO()
        libro.save(salida)
        with self.assertRaises(UserError):
            self._importar(base64.b64encode(salida.getvalue()))

    def test_solo_administrador(self):
        with self.assertRaises(Exception):
            self._importar(self._libro(productos=[('IMP-009', 'X', '', 1, 1, 0, 0)]), user=self.encargado)
